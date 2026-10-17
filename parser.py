"""Reading resumes and pulling candidate details out of them."""
import io
import re
from datetime import date

LEVEL_LABELS = {0: "Not found", 1: "Higher secondary", 2: "Diploma",
                3: "Bachelor's", 4: "Master's", 5: "PhD"}


# ---------------------------------------------------------------- reading files
def extract_text_from_bytes(name: str, data: bytes) -> str:
    """PDF, DOCX or TXT bytes -> plain text."""
    name = name.lower()
    if name.endswith(".pdf"):
        import pdfplumber
        parts = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                parts.append(page.extract_text() or "")
        return "\n".join(parts)

    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:           # many resume templates use tables
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)

    return data.decode("utf-8", errors="ignore")


def extract_text(uploaded_file) -> str:
    """Kept for the old app: accepts a Streamlit uploaded file."""
    return extract_text_from_bytes(uploaded_file.name, uploaded_file.read())


# ---------------------------------------------------------------- contact details
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
YEAR_RANGE_RE = re.compile(r"(19|20)\d{2}\s*[-–—]\s*(19|20)\d{2}")


def find_email(text: str) -> str:
    m = EMAIL_RE.search(text)
    return m.group(0) if m else ""


def find_phone(text: str) -> str:
    for m in re.findall(r"\+?\d[\d\s\-().]{8,16}\d", text):
        digits = re.sub(r"\D", "", m)
        if 10 <= len(digits) <= 13 and not YEAR_RANGE_RE.search(m):
            return m.strip()
    return ""


_NOT_A_NAME = re.compile(
    r"resume|curriculum|vitae|\bcv\b|objective|summary|profile|email|phone|contact|address|"
    r"declaration|experience|education|skills|hotel|manager|intern|trainee|developer|engineer|"
    r"student|fresher|page", re.I)


def guess_name(text: str, filename: str = "") -> str:
    for line in [l.strip() for l in text.splitlines() if l.strip()][:8]:
        words = line.split()
        if (1 <= len(words) <= 4 and re.fullmatch(r"[A-Za-z][A-Za-z.\-' ]+", line)
                and not _NOT_A_NAME.search(line)):
            return line.title() if line.isupper() else line
    base = re.sub(r"\.[A-Za-z0-9]+$", "", filename)
    base = re.sub(r"(?i)resume|cv|curriculum vitae", " ", base)
    base = re.sub(r"[_\-]+", " ", base).strip()
    return base.title() or "Unknown"


# ---------------------------------------------------------------- education
_EDU_LEVELS = [
    (5, r"ph\.?\s?d\b|doctorate|doctoral"),
    (4, r"\bmaster'?s?\s+(?:of|in|degree)|\bmasters\b|\bm\.?\s?tech\b|\bm\.?\s?sc\b|\bmba\b|\bmca\b|"
        r"\bm\.?\s?com\b|\bm\.a\b|\bpgdm\b|post\s?graduate"),
    (3, r"bachelor|\bb\.?\s?tech\b|\bb\.e\b|\bb\.?\s?sc\b|\bbca\b|\bbba\b|\bb\.?\s?com\b|\bb\.a\b|"
        r"\bbhm\w*|undergraduate|graduation|\bdegree\b"),
    (2, r"diploma|polytechnic"),
    (1, r"\bhsc\b|higher secondary|\b12th\b|plus two|\bplus 2\b|pre-?university"),
]


def education_level(text: str):
    """Highest education level found -> (level 0-5, label)."""
    best = 0
    for level, pattern in _EDU_LEVELS:
        if re.search(pattern, text, re.I):
            best = max(best, level)
    return best, LEVEL_LABELS[best]


# ---------------------------------------------------------------- experience
_MONTHS = {m: i + 1 for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
_MON = "|".join(_MONTHS)
_RANGE_RE = re.compile(
    rf"(?:({_MON})[a-z]*\.?\s*,?\s*)?((?:19|20)\d{{2}})\s*(?:-|–|—|to)\s*"
    rf"(?:({_MON})[a-z]*\.?\s*,?\s*)?((?:19|20)\d{{2}}|present|current|now|till date|ongoing)",
    re.I)
_EXP_HEADER = re.compile(
    r"^\W*(?:work(?:ing)?\s+|professional\s+|relevant\s+)?(?:experience|employment|work history|internships?)"
    r"(?:\s+(?:history|details))?\W*$", re.I)
_OTHER_HEADER = re.compile(
    r"^\W*(?:education|academic\w*|skills|technical skills|key skills|projects?|certifications?|"
    r"achievements|languages|hobbies|interests|references?|declaration|summary|objective|"
    r"profile|personal (?:details|information)|training)\W*$", re.I)
_EDU_LINE = re.compile(r"bachelor|b\.?tech|degree|diploma|course|school|college|university|semester|"
                       r"curriculum|master|b\.?sc|m\.?sc", re.I)


def _experience_block(text: str) -> str:
    lines, block, inside = text.splitlines(), [], False
    for line in lines:
        s = line.strip()
        if _EXP_HEADER.match(s):
            inside = True
            continue
        if inside and _OTHER_HEADER.match(s):
            inside = False
        if inside:
            block.append(line)
    return "\n".join(block)


def _months_from_ranges(block: str) -> float:
    now = date.today()
    spans = []
    for m in _RANGE_RE.finditer(block):
        sm, sy, em, ey = m.groups()
        start = int(sy) * 12 + (_MONTHS[sm.lower()] if sm else 1)
        if ey.isdigit():
            end = int(ey) * 12 + (_MONTHS[em.lower()] if em else 1)
        else:
            end = now.year * 12 + now.month
        if end > start:
            spans.append((start, end))
    spans.sort()
    merged = []
    for s, e in spans:                     # merge overlaps so parallel jobs aren't double counted
        if merged and s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    return float(sum(e - s for s, e in merged))


def estimate_experience_years(text: str) -> float:
    """Best guess of work experience: explicit '3 years' phrases or dated jobs."""
    explicit, months_only = [], []
    for line in text.splitlines():
        if _EDU_LINE.search(line):         # '4 years B.Tech' is not work experience
            continue
        for n in re.findall(r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)", line, re.I):
            explicit.append(float(n))
        for n in re.findall(r"(\d{1,2})\s*\+?\s*months?", line, re.I):
            months_only.append(int(n) / 12)
    from_text = max(explicit) if explicit else (max(months_only) if months_only else 0.0)
    from_dates = _months_from_ranges(_experience_block(text)) / 12
    return round(min(max(from_text, from_dates), 40.0), 1)


def required_years(job_text: str) -> float:
    """Years of experience a job description asks for (0 if none)."""
    nums = [int(n) for n in re.findall(
        r"(\d{1,2})(?:\s*(?:-|–|to)\s*\d{1,2})?\s*\+?\s*(?:years?|yrs?)", job_text, re.I)]
    nums = [n for n in nums if n <= 15]    # ignore things like '25 years in business'
    return float(max(nums)) if nums else 0.0


# ---------------------------------------------------------------- everything at once
def candidate_info(text: str, filename: str = "") -> dict:
    level, label = education_level(text)
    return {
        "name": guess_name(text, filename),
        "email": find_email(text),
        "phone": find_phone(text),
        "edu_level": level,
        "education": label,
        "years": estimate_experience_years(text),
    }
